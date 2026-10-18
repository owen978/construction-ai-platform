"""Schedule the daily-ideas pipeline (365 days) and write the review workbook."""
import datetime as dt
import json
import sys
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, sys.argv[1])
from pipeline_ideas import MON, TUE, WED, THU, FRI, SAT, SUN  # noqa: E402

OUT = sys.argv[2]
BACKLOG = json.loads(open(sys.argv[3]).read())    # [{n, kw, slug}] in engine order
LIVE = json.loads(open(sys.argv[4]).read())       # {'workflows': [...], 'guides': [...], 'templates': [...]}

START = dt.date(2026, 10, 26)
DAYS = 365
THEMES = {
    0: ('Commercial & QS', MON), 1: ('Site, quality & handover', TUE), 2: ('Projects, programme & design', WED),
    3: ('Contracts, claims & compliance', THU), 4: ('Estimating, tendering & procurement', FRI),
    5: ('Health, safety & environment', SAT), 6: ('AI tools, agents & the business', SUN),
}
DIFF = {'b': 'Beginner', 'i': 'Intermediate', 'a': 'Advanced'}
RISK = {'l': 'Low', 'm': 'Medium', 'h': 'High'}
ROLE = {
    'quantity-surveyor': 'QS', 'commercial-manager': 'Commercial Manager', 'project-manager': 'Project Manager',
    'site-manager': 'Site Manager', 'contracts-manager': 'Contracts Manager', 'construction-director': 'Director',
    'health-and-safety-manager': 'H&S Manager', 'estimator': 'Estimator', 'planning-engineer': 'Planner',
    'design-manager': 'Design Manager', 'bim-manager': 'BIM Manager', 'procurement-manager': 'Buyer',
    'document-control': 'Document Controller', 'subcontractor': 'Subcontractor',
}
TOOL = {'claude': 'Claude', 'chatgpt': 'ChatGPT', 'copilot': 'Microsoft Copilot', 'claude-code': 'Claude Code',
        'gemini': 'Gemini', 'grok': 'Grok', 'any': 'Any', 'gamma': 'Gamma', 'perplexity': 'Perplexity',
        'notion': 'Notion AI', 'bluebeam': 'Bluebeam'}
LINKEDIN = {  # Owen's scheduled LinkedIn posts (Metricool, drafts at 09:15)
    dt.date(2026, 10, 21): 'Quote comparison: the £1,100 access exclusion',
    dt.date(2026, 10, 28): 'Trial one reporting task properly',
    dt.date(2026, 11, 4): 'Invoice with no project reference: what should the workflow do?',
    dt.date(2026, 11, 11): 'Meeting actions need an owner and a date',
    dt.date(2026, 11, 18): "Don't make a chat response the cash-flow model",
    dt.date(2026, 11, 25): 'Friday updates from four places',
    dt.date(2026, 12, 2): 'Who owns the outstanding records at closeout?',
    dt.date(2026, 12, 9): 'Bring me one recurring task',
}

# Backlog guides: the weekly content task publishes the next queued row every Monday and Thursday.
publish_days = []
d = dt.date(2026, 10, 19)
while len(publish_days) < len(BACKLOG):
    if d.weekday() in (0, 3):
        publish_days.append(d)
    d += dt.timedelta(days=1)
GUIDE_ON = {day: row for day, row in zip(publish_days, BACKLOG)}

# Assign pinned ideas first, then fill each weekday in list order.
pinned, queues = {}, {}
for wd, (_, items) in THEMES.items():
    queues[wd] = []
    for it in items:
        if it['pin']:
            p = dt.date.fromisoformat(it['pin'])
            assert p.weekday() == wd, f"{it['title']} pinned to {p} ({p:%A}) but belongs to weekday {wd}"
            assert p not in pinned, f'two ideas pinned to {p}'
            pinned[p] = it
        else:
            queues[wd].append(it)

rows = []
for n in range(DAYS):
    day = START + dt.timedelta(days=n)
    wd = day.weekday()
    it = pinned.pop(day, None) or queues[wd].pop(0)
    rows.append((n + 1, day, THEMES[wd][0], it))
assert not pinned, f'unused pins: {list(pinned)}'
reserve = [(THEMES[wd][0], it) for wd in queues for it in queues[wd]]

titles = [r[3]['title'] for r in rows] + [it['title'] for _, it in reserve]
dupes = [t for t, c in Counter(titles).items() if c > 1]
assert not dupes, f'duplicate titles: {dupes}'

# ── Workbook ─────────────────────────────────────────────────────────
wb = Workbook()
HEAD = Font(bold=True, color='FFFFFF')
FILL = PatternFill('solid', fgColor='1A1A2E')
WRAP = Alignment(wrap_text=True, vertical='top')

def sheet(ws, headers, widths):
    ws.append(headers)
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
        cell = ws.cell(row=1, column=c)
        cell.font, cell.fill, cell.alignment = HEAD, FILL, Alignment(wrap_text=True, vertical='center')
    ws.freeze_panes = 'A2'

ws = wb.active
ws.title = 'Pipeline (365 days)'
headers = ['Day', 'Date', 'Weekday', 'Theme', 'Idea', 'Roles', 'Task', 'Contract form', 'Tool', 'Difficulty',
           'Risk if wrong', 'Builds on (live page)', 'Source', 'Notes', 'Guide the content task publishes that day',
           'Your LinkedIn post that day', 'Approve? (Yes / No / Change)', 'Your notes']
sheet(ws, headers, [5, 11, 10, 24, 52, 26, 18, 16, 14, 12, 11, 34, 24, 40, 30, 30, 14, 36])
for n, day, theme, it in rows:
    g = GUIDE_ON.get(day)
    ws.append([
        n, day, f'{day:%A}', theme, it['title'], ', '.join(ROLE.get(r, r) for r in it['roles']), it['task'],
        it['contract'], TOOL.get(it['tool'], it['tool']), DIFF[it['difficulty']], RISK[it['risk']],
        ('https://buildcopilot.ai' + it['related']) if it['related'] else '', it['source'], it['note'],
        f"{g['kw']} (/guides/{g['slug']})" if g else '', LINKEDIN.get(day, ''), '', '',
    ])
for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = WRAP
    row[1].number_format = 'ddd d mmm yyyy'
ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{len(rows) + 1}'
dv = DataValidation(type='list', formula1='"Yes,No,Change"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f'Q2:Q{len(rows) + 1}')

rs = wb.create_sheet('Reserve')
sheet(rs, ['Theme', 'Idea', 'Roles', 'Task', 'Tool', 'Notes', 'Use instead of (day #)'], [24, 52, 26, 18, 14, 40, 18])
for theme, it in reserve:
    rs.append([theme, it['title'], ', '.join(ROLE.get(r, r) for r in it['roles']), it['task'],
               TOOL.get(it['tool'], it['tool']), it['note'], ''])

cv = wb.create_sheet('Coverage')
sheet(cv, ['Measure', 'Value', 'Days'], [20, 40, 8])
for label, counter in [
    ('Theme', Counter(r[2] for r in rows)),
    ('Role', Counter(ROLE.get(x, x) for r in rows for x in r[3]['roles'])),
    ('Tool', Counter(TOOL.get(r[3]['tool'], r[3]['tool']) for r in rows)),
    ('Difficulty', Counter(DIFF[r[3]['difficulty']] for r in rows)),
    ('Risk if wrong', Counter(RISK[r[3]['risk']] for r in rows)),
    ('Contract form', Counter(c.strip() for r in rows for c in (r[3]['contract'] or 'none').split(','))),
    ('Task', Counter(r[3]['task'] for r in rows)),
]:
    for value, n in counter.most_common():
        cv.append([label, value, n])

ex = wb.create_sheet('Already on the site')
sheet(ex, ['Type', 'Title', 'URL'], [12, 70, 60])
for kind, path in [('workflows', 'ai-workflows'), ('guides', 'guides'), ('templates', 'templates')]:
    for item in LIVE[kind]:
        ex.append([kind[:-1].title(), item['title'], f"https://buildcopilot.ai/{path}/{item['slug']}"])

hw = wb.create_sheet('How to review')
hw.column_dimensions['A'].width = 120
for line in [
    'BuildCopilot daily ideas: 365-day pipeline for review (Mon 26 Oct 2026 to Mon 25 Oct 2027).',
    '',
    'Each weekday has a theme, so every role and use case comes round every week:',
    '  Mon Commercial & QS · Tue Site, quality & handover · Wed Projects, programme & design · Thu Contracts, claims & compliance',
    '  Fri Estimating, tendering & procurement · Sat Health, safety & environment (incl. the toolbox talk series) · Sun AI tools, agents & the business',
    '',
    'How to review: put Yes, No or Change in the Approve column (it has a dropdown). For Change, say what you want in Your notes.',
    'Anything marked No gets replaced from the Reserve sheet or a new idea of the same theme. Approved rows become the daily pipeline.',
    '',
    'Left out on purpose: anything that already has its own workflow, guide or template page (see "Already on the site").',
    'Included but flagged in Notes: ideas that already appear as a short prompt inside a published guide. Untick them if you only want brand-new topics.',
    'Seasonal ideas are pinned to their dates (dark mornings after the clocks change, cold weather, Christmas shutdown and restart, March year end,',
    '  Mental Health Awareness Week, summer heat). A few ideas sit next to your scheduled LinkedIn posts so the topics reinforce each other.',
    'The "Guide the content task publishes" column shows the queued backlog guide your Mon/Thu content task is due to publish that day.',
    '',
    'Time saved is not in this sheet: each idea gets cautious "typical" by-hand and with-AI times when it is written.',
    'Contract notices (payment, pay less, NEC4, JCT, adjudication, suspension) carry a "check your contract and get advice" line and a High risk flag.',
    'Ideas drawn from real client work (Bluefield) are marked "Anonymise" and will use sizes and trades only, never names.',
]:
    hw.append([line])

wb.move_sheet('How to review', offset=-4)
wb.active = 0
wb.save(OUT)

print(f'scheduled {len(rows)} days ({rows[0][1]} to {rows[-1][1]}), reserve {len(reserve)}')
print('themes:', dict(Counter(r[2] for r in rows)))
print('roles:', dict(Counter(ROLE.get(x, x) for r in rows for x in r[3]['roles']).most_common()))
print('flagged as already a short prompt in a guide:', sum('Already a short prompt' in r[3]['note'] for r in rows))
print('backlog guides on the calendar:', len(GUIDE_ON), 'through', publish_days[-1])
